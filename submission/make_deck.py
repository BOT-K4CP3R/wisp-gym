#!/usr/bin/env python3
"""Builds submission/wisp-gym-presentation.pptx (16:9, dark) from the real emulator captures.
Needs: python-pptx, pillow."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.abspath(__file__))
A = lambda *p: os.path.join(ROOT, 'assets', *p)
OUT = os.path.join(ROOT, 'wisp-gym-presentation.pptx')

BG, CARD, EDGE = RGBColor(10, 10, 12), RGBColor(21, 22, 25), RGBColor(48, 49, 54)
ACC, TXT, DIM, MUTE = RGBColor(255, 138, 91), RGBColor(244, 244, 245), RGBColor(161, 161, 170), RGBColor(113, 113, 122)
FONT = 'Arial'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    return s

def text(s, x, y, w, h, t, size=18, color=TXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
    return tb

def card(s, x, y, w, h, radius=0.08):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid(); sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = EDGE; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh

def accent_bar(s, x, y):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.7), Inches(0.08))
    b.adjustments[0] = 0.5; b.fill.solid(); b.fill.fore_color.rgb = ACC; b.line.fill.background()

def title(s, head, sub=None, w=11.5):
    accent_bar(s, 0.7, 0.62)
    text(s, 0.7, 0.8, w, 1.0, head, size=36, bold=True)
    if sub:
        text(s, 0.7, 1.62, w, 0.6, sub, size=16, color=DIM)

def phone(s, path, x, y, h=5.6):
    """Phone screenshot with a rounded bezel."""
    from PIL import Image
    w = h * 360 / 720
    bez = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - 0.1), Inches(y - 0.1), Inches(w + 0.2), Inches(h + 0.2))
    bez.adjustments[0] = 0.09; bez.fill.solid(); bez.fill.fore_color.rgb = RGBColor(34, 35, 40)
    bez.line.color.rgb = RGBColor(70, 71, 78); bez.line.width = Pt(1.2)
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return pic

def footer(s, n):
    text(s, 0.7, 7.05, 8, 0.3, 'Wisp Gym  |  HackYeah 2026  |  Huawei Challenge: Imagine What\'s Next', size=10, color=MUTE)
    text(s, 12.0, 7.05, 0.7, 0.3, str(n), size=10, color=MUTE, align=PP_ALIGN.RIGHT)

def bullets(s, x, y, w, items, size=18, gap=0.78):
    for i, (head, body) in enumerate(items):
        yy = y + i * gap
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.1), Inches(0.14), Inches(0.14))
        dot.fill.solid(); dot.fill.fore_color.rgb = ACC; dot.line.fill.background()
        text(s, x + 0.35, yy, w - 0.35, 0.4, head, size=size, bold=True)
        if body:
            text(s, x + 0.35, yy + 0.36, w - 0.35, 0.5, body, size=size - 4, color=DIM)

n = 0
# 1 --- title
s = slide(); n += 1
s.shapes.add_picture(A('icon.png'), Inches(5.67), Inches(1.0), Inches(2.0), Inches(2.0))
text(s, 0, 3.35, 13.333, 1.1, 'Wisp Gym', size=60, bold=True, align=PP_ALIGN.CENTER)
text(s, 0, 4.45, 13.333, 0.6, 'A creature that grows from your real training', size=24, color=DIM, align=PP_ALIGN.CENTER)
text(s, 0, 5.6, 13.333, 0.4, 'Native OpenHarmony / HarmonyOS app  |  ArkTS + ArkUI  |  API 20+', size=16, color=MUTE, align=PP_ALIGN.CENTER)
text(s, 0, 6.2, 13.333, 0.4, 'HackYeah 2026  |  Huawei Challenge: Imagine What\'s Next  |  Human-Centric Technology + Spatial Experiences', size=13, color=MUTE, align=PP_ALIGN.CENTER)

# 2 --- problem
s = slide(); n += 1
title(s, 'Training apps get the basics wrong', 'The friction is in the middle of the set, and the pressure is in the streak.')
cards = [('Manual logging', 'You count and type reps with sweaty hands, or you guess.'),
         ('No set feedback', 'Nothing tells you when a set is getting too slow to be useful.'),
         ('Pressure, not balance', 'Streaks and guilt push people to train through fatigue instead of recovering.')]
for i, (h, b) in enumerate(cards):
    x = 0.7 + i * 4.1
    card(s, x, 2.6, 3.85, 3.2)
    text(s, x + 0.35, 3.0, 3.2, 0.6, f'0{i+1}', size=30, bold=True, color=ACC)
    text(s, x + 0.35, 3.8, 3.2, 0.6, h, size=22, bold=True)
    text(s, x + 0.35, 4.5, 3.2, 1.2, b, size=15, color=DIM)
footer(s, n)

# 3 --- solution
s = slide(); n += 1
title(s, 'Your training, made visible', w=7.4)
phone(s, A('home_thriving.jpg'), 8.6, 1.0, h=5.7)
bullets(s, 0.7, 2.0, 7.4, [
    ('Reps counted by the phone', 'Accelerometer only, any orientation in the pocket.'),
    ('Knows when to stop', 'Tempo and drive loss trigger a "rack it" cue, with the numbers.'),
    ('Rest is part of the game', 'Over-training makes Wisp ask for a rest day. It never dies.'),
    ('Explainable and private', 'Every value is traceable. Everything stays on the phone.')], size=20, gap=1.15)
footer(s, n)

# 4 --- how it works
s = slide(); n += 1
title(s, 'How it works', 'Pure, tested logic in the middle; thin OS adapters at the edges.')
steps = [('Accelerometer', 'SensorServiceKit, 50 Hz'), ('Rep detector', 'gravity removal, peak detection, walking rejection'), ('Set analyzer', 'tempo + drive loss -> fatigue, stop cue'),
         ('Life engine', 'decay, recovery, rest guard, reasons'), ('Creature + widget', 'procedural Canvas, service widget')]
for i, (h, b) in enumerate(steps):
    x = 0.7 + i * 2.5
    card(s, x, 2.9, 2.2, 2.5)
    text(s, x + 0.2, 3.15, 1.9, 0.8, h, size=17, bold=True)
    text(s, x + 0.2, 4.0, 1.85, 1.2, b, size=12, color=DIM)
    if i < len(steps) - 1:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.22), Inches(4.0), Inches(0.26), Inches(0.28))
        a.fill.solid(); a.fill.fore_color.rgb = ACC; a.line.fill.background()
text(s, 0.7, 5.9, 12, 0.9, ['No AI model: signal processing and transparent rules.', 'The same source files are unit-tested on Node and compiled for the device.'], size=15, color=DIM)
footer(s, n)

# 5 --- live demo frames
s = slide(); n += 1
title(s, 'A real set on the emulator', 'Reps climb as the ring fills; at 25% fatigue the ring turns orange and Wisp says "rack it".')
for i, f in enumerate(['04.jpg', '08.jpg', '16.jpg']):
    phone(s, A('set', f), 1.55 + i * 3.9, 2.3, h=4.4)
text(s, 0.7, 6.8, 12, 0.3, 'Emulator capture, Demo mode: a simulated accelerometer stream replayed through the same detector as the real sensor.', size=11, color=MUTE)
footer(s, n)

# 6 --- design principles
s = slide(); n += 1
title(s, 'Motivate, don\'t pressure', w=6.0)
phone(s, A('home_rest.jpg'), 6.9, 1.1, h=5.5)
phone(s, A('insights.jpg'), 9.9, 1.1, h=5.5)
bullets(s, 0.7, 2.0, 6.0, [
    ('Non-punitive by construction', 'Levels have a floor. No streaks to lose. Never sad.'),
    ('Recovery as a mechanic', 'Three sessions in 72 h and Wisp asks to rest.'),
    ('Explainable', 'The Insights screen shows the numbers behind every value.')], size=19, gap=1.3)
footer(s, n)

# 7 --- platform
s = slide(); n += 1
title(s, 'What it uses from the platform')
rows = [('Accelerometer', 'SensorServiceKit', 'live rep detection'), ('Pedometer', 'SensorServiceKit + ACTIVITY_MOTION', 'daily movement for Wisp'),
        ('Service widget', 'FormExtensionAbility, formProvider', 'creature + status on the home screen'), ('Vibrator', 'SensorServiceKit', 'stop cue, end of rest'),
        ('Notifications', 'NotificationKit', 'rest finished'), ('Preferences', '@ohos.data.preferences', 'private, validated local storage'),
        ('Window', 'keepScreenOn, system bars', 'workout screen stays on')]
for i, (a, b, c) in enumerate(rows):
    y = 2.0 + i * 0.68
    card(s, 0.7, y, 11.9, 0.58, radius=0.25)
    text(s, 1.0, y + 0.13, 3.0, 0.4, a, size=16, bold=True)
    text(s, 4.2, y + 0.15, 4.4, 0.4, b, size=13, color=DIM)
    text(s, 8.7, y + 0.15, 3.8, 0.4, c, size=13, color=ACC)
footer(s, n)

# 8 --- quality
s = slide(); n += 1
title(s, 'Built to be checked', 'Claims are backed by tests, builds and emulator runs.')
items = [('22', 'unit tests on the real source files'), ('300', 'randomised synthetic sets: 99.3% exact, 100% within 1 rep'), ('0', 'network permissions, secrets or accounts'), ('API 20', 'minimum; compiled on 23, runs on the OpenHarmony 6.1 emulator')]
for i, (big, small) in enumerate(items):
    x = 0.7 + (i % 2) * 6.1; y = 2.5 + (i // 2) * 2.1
    card(s, x, y, 5.8, 1.85)
    text(s, x + 0.4, y + 0.3, 2.2, 1.0, big, size=44, bold=True, color=ACC)
    text(s, x + 2.7, y + 0.5, 2.9, 1.2, small, size=15, color=DIM)
text(s, 0.7, 6.7, 12, 0.3, 'Reproducible: ./scripts/setup-macos.sh, test.sh, build.sh, emulator-up.sh. DevEco Studio is not required.', size=12, color=MUTE)
footer(s, n)

# 9 --- honest status
s = slide(); n += 1
title(s, 'What is real, and what is not', 'We say it plainly.')
card(s, 0.7, 2.3, 5.9, 4.3); card(s, 6.8, 2.3, 5.8, 4.3)
text(s, 1.1, 2.6, 5.2, 0.5, 'Verified', size=22, bold=True, color=ACC)
text(s, 1.1, 3.3, 5.2, 3.2, ['Builds and signs; installs and launches on an OpenHarmony 6.1 emulator', 'All screens, a live-counted set with the stop cue, rest timer', 'Home-screen widget added and updated', 'System permission dialog with a stated reason', '22 unit tests'], size=14, color=TXT, spacing=1.25)
text(s, 7.2, 2.6, 5.0, 0.5, 'Not yet', size=22, bold=True, color=ACC)
text(s, 7.2, 3.3, 5.0, 3.2, ['Real accelerometer and pedometer data (emulators have no motion): the demo replays a labelled simulated stream', 'Rep accuracy on real bodies: tested on synthetic traces only', 'Widget tap-through; HarmonyOS device; DevEco emulator'], size=14, color=TXT, spacing=1.25)
footer(s, n)

# 10 --- AI workflow
s = slide(); n += 1
title(s, 'How AI was used', 'The product contains no AI model. The development did.')
bullets(s, 0.7, 2.4, 11.5, [
    ('Claude Sonnet 5.5 in Claude Code', 'Ideation, architecture, code, tests, docs, emulator-driven debugging. Documented in AI_WORKFLOW.md.'),
    ('Validated, not trusted', 'Compiler, 22 tests, on-device screenshots and logs. Bugs found on the emulator are listed with their fixes.'),
    ('Human decisions', 'Concept, scope (no AI in the product, an app not a feature), non-punitive design, visual direction.')], size=20, gap=1.45)
footer(s, n)

# 11 --- roadmap / close
s = slide(); n += 1
title(s, 'Next')
bullets(s, 0.7, 2.0, 6.5, [
    ('Validate on real recordings', 'Squats, deadlifts, curls on a real body, in a pocket or armband.'),
    ('Watch profile', 'Wrist sensor and heart rate for the same creature.'),
    ('More exercises, social creatures', 'QR-based creature exchange, still without a server.')], size=19, gap=1.4)
phone(s, A('widget.jpg'), 9.2, 1.1, h=5.5)
text(s, 0.7, 6.3, 7, 0.5, 'Thank you', size=30, bold=True, color=ACC)
footer(s, n)

prs.save(OUT)
print('saved', OUT, os.path.getsize(OUT) // 1024, 'KB')
